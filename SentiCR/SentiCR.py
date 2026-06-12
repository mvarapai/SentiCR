from __future__ import annotations

import argparse
import csv
import random
import re
from dataclasses import dataclass
from importlib.resources import as_file, files
from statistics import mean
from typing import Iterable

import nltk
import numpy as np
from imblearn.over_sampling import SMOTE
from nltk.stem.snowball import SnowballStemmer
from openpyxl import load_workbook
from sklearn.ensemble import AdaBoostClassifier
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import SGDClassifier
from sklearn.metrics import accuracy_score
from sklearn.metrics import f1_score
from sklearn.metrics import precision_score
from sklearn.metrics import recall_score
from sklearn.model_selection import KFold
from sklearn.naive_bayes import BernoulliNB
from sklearn.neural_network import MLPClassifier
from sklearn.svm import LinearSVC
from sklearn.tree import DecisionTreeClassifier


PACKAGE_FILES = files("SentiCR")
CONTRACTIONS_PATH = PACKAGE_FILES / "Contractions.txt"
EMOTICONS_PATH = PACKAGE_FILES / "EmoticonLookupTable.txt"
ORACLE_PATH = PACKAGE_FILES / "oracle.xlsx"


def replace_all(text: str, replacements: dict[str, str]) -> str:
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


stemmer = SnowballStemmer("english")


def stem_tokens(tokens: Iterable[str]) -> list[str]:
    return [stemmer.stem(token) for token in tokens]


def tokenize_and_stem(text: str) -> list[str]:
    tokens = nltk.word_tokenize(text)
    return stem_tokens(tokens)


mystop_words = [
    "i", "me", "my", "myself", "we", "our", "ourselves", "you", "your",
    "yourself", "yourselves", "he", "him", "his", "himself", "she", "her",
    "herself", "it", "its", "itself", "they", "them", "their", "themselves",
    "this", "that", "these", "those", "am", "is", "are", "was", "were",
    "be", "been", "being", "have", "has", "had", "having", "do", "does",
    "did", "doing", "a", "an", "the", "and", "if", "or", "as", "until",
    "of", "at", "by", "between", "into", "through", "during", "to", "from",
    "in", "out", "on", "off", "then", "once", "here", "there", "all", "any",
    "both", "each", "few", "more", "other", "some", "such", "than", "too",
    "very", "s", "t", "can", "will", "don", "should", "now",

    # Programming keywords
    "while", "case", "switch", "def", "abstract", "byte", "continue",
    "native", "private", "synchronized", "include", "finally", "class",
    "double", "float", "int", "else", "instanceof", "long", "super",
    "import", "short", "default", "catch", "try", "new", "final", "extends",
    "implements", "public", "protected", "static", "return", "char", "const",
    "break", "boolean", "bool", "package", "assert", "raise", "global",
    "with", "yield", "except", "enum", "signed", "void", "virtual", "union",
    "goto", "var", "function", "require", "print", "echo", "foreach",
    "elseif", "namespace", "delegate", "event", "override", "struct",
    "readonly", "explicit", "interface", "get", "set", "elif", "for",
    "throw", "throws", "lambda", "endfor", "endforeach", "endif", "endwhile",
    "clone",
]


def load_dictionary(path, delimiter: str = "\t") -> dict[str, str]:
    with path.open("r", encoding="utf-8") as file:
        reader = csv.reader(file, delimiter=delimiter)
        return {row[0]: row[1] for row in reader if len(row) >= 2}


contractions_dict = load_dictionary(CONTRACTIONS_PATH)
emodict = load_dictionary(EMOTICONS_PATH)


grammar = r"""
NegP: {<VERB>?<ADV>+<VERB|ADJ>?<PRT|ADV><VERB>}
      {<VERB>?<ADV>+<VERB|ADJ>*<ADP|DET>?<ADJ>?<NOUN>?<ADV>?}
"""

chunk_parser = nltk.RegexpParser(grammar)

contractions_regex = re.compile(
    r"(%s)" % "|".join(re.escape(key) for key in contractions_dict.keys())
)

url_regex = re.compile(
    r"http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\(\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+"
)

negation_words = [
    "not", "never", "none", "nobody", "nowhere", "neither", "barely",
    "hardly", "nothing", "rarely", "seldom", "despite",
]

emoticon_words = ["PositiveSentiment", "NegativeSentiment"]


def expand_contractions(text: str) -> str:
    def replace(match: re.Match[str]) -> str:
        return contractions_dict[match.group(0)]

    return contractions_regex.sub(replace, text.lower())


def remove_url(text: str) -> str:
    return url_regex.sub(" ", text)


def negated(input_words: Iterable[str]) -> bool:
    return any(word in input_words for word in negation_words)


def prepend_not(word: str) -> str:
    if word in emoticon_words:
        return word
    if word in negation_words:
        return word
    return "NOT_" + word


def handle_negation(comments: str) -> str:
    sentences = nltk.sent_tokenize(comments)
    modified_sentences = []

    for sentence in sentences:
        words = nltk.word_tokenize(sentence)

        if not negated(words):
            modified_sentences.append(sentence)
            continue

        tagged_words = nltk.tag.pos_tag(words, tagset="universal")
        chunked = chunk_parser.parse(tagged_words)
        modified_words = []

        for node in chunked:
            if isinstance(node, nltk.tree.Tree):
                chunk_words = [pair[0] for pair in node.leaves()]

                if node.label() == "NegP" and negated(chunk_words):
                    for word, pos in node.leaves():
                        if pos in {"ADV", "ADJ", "VERB"} and word != "not":
                            modified_words.append(prepend_not(word))
                        else:
                            modified_words.append(word)
                else:
                    modified_words.extend(chunk_words)
            else:
                modified_words.append(node[0])

        modified_sentences.append(" ".join(modified_words))

    return ". ".join(modified_sentences)


def preprocess_text(text: object) -> str:
    if text is None:
        comments = ""
    elif isinstance(text, bytes):
        comments = text.decode("utf-8", errors="ignore")
    else:
        comments = str(text)

    comments = expand_contractions(comments)
    comments = remove_url(comments)
    comments = replace_all(comments, emodict)
    comments = handle_negation(comments)

    return comments


@dataclass
class SentimentData:
    text: str
    rating: int


def read_oracle_data() -> list[SentimentData]:
    with as_file(ORACLE_PATH) as oracle_path:
        wb = load_workbook(oracle_path, data_only=True)

    ws = wb.worksheets[0]
    oracle_data = []

    for row in ws.iter_rows(values_only=True):
        if not row or (row[0] is None and row[1] is None):
            continue

        text, rating = row[0], row[1]

        if isinstance(text, str) and text.strip().lower() in {
            "text", "comment", "comments",
        }:
            continue

        oracle_data.append(SentimentData(str(text), int(rating)))

    return oracle_data


class SentiCR:
    def __init__(self, algo: str = "GBT", training_data: Iterable[SentimentData] | None = None):
        self.algo = algo

        if training_data is None:
            print("Reading data from oracle..")
            self.training_data = read_oracle_data()
        else:
            self.training_data = list(training_data)

        self.vectorizer: TfidfVectorizer | None = None
        self.model = self.create_model_from_training_data()

    def get_classifier(self):
        if self.algo == "GBT":
            return GradientBoostingClassifier()
        if self.algo == "RF":
            return RandomForestClassifier()
        if self.algo == "ADB":
            return AdaBoostClassifier()
        if self.algo == "DT":
            return DecisionTreeClassifier()
        if self.algo == "NB":
            return BernoulliNB()
        if self.algo == "SGD":
            return SGDClassifier()
        if self.algo == "SVC":
            return LinearSVC()
        if self.algo == "MLPC":
            return MLPClassifier(
                activation="logistic",
                batch_size="auto",
                early_stopping=True,
                hidden_layer_sizes=(100,),
                learning_rate="adaptive",
                learning_rate_init=0.1,
                max_iter=5000,
                random_state=1,
                solver="lbfgs",
                tol=0.0001,
                validation_fraction=0.1,
                verbose=False,
                warm_start=False,
            )

        raise ValueError(f"Unknown classifier algorithm: {self.algo}")

    def create_model_from_training_data(self):
        training_comments = []
        training_ratings = []

        print("Training classifier model..")

        for item in self.training_data:
            training_comments.append(preprocess_text(item.text))
            training_ratings.append(item.rating)

        self.vectorizer = TfidfVectorizer(
            tokenizer=tokenize_and_stem,
            token_pattern=None,
            sublinear_tf=True,
            max_df=0.5,
            stop_words=mystop_words,
            min_df=3,
        )

        x_train = self.vectorizer.fit_transform(training_comments).toarray()
        y_train = np.array(training_ratings)

        smote_model = SMOTE(sampling_strategy=0.5, k_neighbors=5)
        x_resampled, y_resampled = smote_model.fit_resample(x_train, y_train)

        model = self.get_classifier()
        model.fit(x_resampled, y_resampled)

        return model

    def get_sentiment_polarity(self, text: object) -> int:
        if self.vectorizer is None:
            raise RuntimeError("Vectorizer has not been initialized.")

        comment = preprocess_text(text)
        feature_vector = self.vectorizer.transform([comment]).toarray()
        return int(self.model.predict(feature_vector)[0])

    def get_sentiment_polarity_collection(self, texts: Iterable[object]) -> list[int]:
        return [self.get_sentiment_polarity(text) for text in texts]


def ten_fold_cross_validation(dataset: np.ndarray, algo: str, random_state: int | None = None):
    kf = KFold(n_splits=10, shuffle=True, random_state=random_state)

    run_precision = []
    run_recall = []
    run_f1score = []
    run_accuracy = []

    for count, (train, test) in enumerate(kf.split(dataset), start=1):
        print("Using split-" + str(count) + " as test data..")

        classifier_model = SentiCR(
            algo=algo,
            training_data=dataset[train],
        )

        test_comments = [item.text for item in dataset[test]]
        test_ratings = [item.rating for item in dataset[test]]

        predictions = classifier_model.get_sentiment_polarity_collection(test_comments)

        precision = precision_score(test_ratings, predictions, pos_label=-1)
        recall = recall_score(test_ratings, predictions, pos_label=-1)
        f1score = f1_score(test_ratings, predictions, pos_label=-1)
        accuracy = accuracy_score(test_ratings, predictions)

        run_precision.append(precision)
        run_recall.append(recall)
        run_f1score.append(f1score)
        run_accuracy.append(accuracy)

    return (
        mean(run_precision),
        mean(run_recall),
        mean(run_f1score),
        mean(run_accuracy),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Supervised sentiment classifier")

    parser.add_argument(
        "--algo",
        type=str,
        default="GBT",
        help="Classification algorithm",
    )

    parser.add_argument(
        "--repeat",
        type=int,
        default=100,
        help="Iteration count",
    )

    args = parser.parse_args()

    algo = args.algo
    repeat = args.repeat

    print("Cross validation")
    print("Algorithm: " + algo)
    print("Repeat: " + str(repeat))

    oracle_data = read_oracle_data()
    random.shuffle(oracle_data)
    oracle_data_array = np.array(oracle_data, dtype=object)

    precision_runs = []
    recall_runs = []
    fmean_runs = []
    accuracy_runs = []

    for run_index in range(repeat):
        print(".............................")
        print("Run# {}".format(run_index))

        precision, recall, f1score, accuracy = ten_fold_cross_validation(
            oracle_data_array,
            algo,
            random_state=run_index,
        )

        precision_runs.append(precision)
        recall_runs.append(recall)
        fmean_runs.append(f1score)
        accuracy_runs.append(accuracy)

        print("Precision:" + str(precision))
        print("Recall:" + str(recall))
        print("F-measure:" + str(f1score))
        print("Accuracy:" + str(accuracy))

    output_file = "cross-validation-" + algo + ".csv"

    with open(output_file, "w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Run", "Algo", "Precision", "Recall", "Fscore", "Accuracy"])

        for run_index in range(repeat):
            writer.writerow([
                run_index,
                algo,
                precision_runs[run_index],
                recall_runs[run_index],
                fmean_runs[run_index],
                accuracy_runs[run_index],
            ])

    print("-------------------------")
    print("Average Precision: {}".format(mean(precision_runs)))
    print("Average Recall: {}".format(mean(recall_runs)))
    print("Average Fmean: {}".format(mean(fmean_runs)))
    print("Average Accuracy: {}".format(mean(accuracy_runs)))
    print("-------------------------")


if __name__ == "__main__":
    main()